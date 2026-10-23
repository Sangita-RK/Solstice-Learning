import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from main import answer_question
from docs_qa_utils import load_docs_from_folder


ROOT = Path(__file__).resolve().parents[1]


def write_text_pdf(path, text):
    encoded_text = text.encode("ascii")
    stream = b"BT /F1 12 Tf 72 720 Td (" + encoded_text + b") Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        (
            b"<< /Length "
            + str(len(stream)).encode("ascii")
            + b" >>\nstream\n"
            + stream
            + b"\nendstream"
        ),
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("ascii"))
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")

    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        (
            f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    path.write_bytes(pdf)


class AnswerQuestionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.docs = load_docs_from_folder(ROOT / "data" / "docs")
        with (ROOT / "data" / "example_questions.json").open(encoding="utf-8") as file:
            cls.examples = json.load(file)["example_questions"]

    def test_provided_supported_questions_return_source_passages(self):
        for example in self.examples:
            if example["expected_doc_id"] is None:
                continue
            with self.subTest(question=example["question"]):
                result = answer_question(example["question"], self.docs)
                self.assertEqual(result.status, "answer")
                self.assertEqual(result.doc_id, example["expected_doc_id"])
                source_doc = next(
                    doc for doc in self.docs if doc["doc_id"] == result.doc_id
                )
                self.assertIn(result.passage, source_doc["content"])

    def test_paraphrased_question_returns_supporting_passage(self):
        result = answer_question("What is the CSV export row limit?", self.docs)

        self.assertEqual(result.status, "answer")
        self.assertEqual(result.doc_id, "D-03")
        self.assertIn("10,000 rows", result.passage)

    def test_password_change_paraphrase_returns_reset_instructions(self):
        for question in (
            "How do I change my password?",
            "I need to update my password",
            "changing password",
            "changing my password",
        ):
            with self.subTest(question=question):
                result = answer_question(question, self.docs)

                self.assertEqual(result.status, "answer")
                self.assertEqual(result.doc_id, "D-01")
                self.assertIn("Forgot password", result.passage)
                self.assertIn("reset link", result.passage)

    def test_salesforce_question_abstains(self):
        example = next(
            item for item in self.examples if item["expected_doc_id"] is None
        )

        result = answer_question(example["question"], self.docs)

        self.assertEqual(result.status, "abstain")
        self.assertIsNone(result.passage)
        self.assertEqual(result.message, "I don't have an answer for that.")

    def test_unmatched_specific_product_prevents_partial_answer(self):
        result = answer_question(
            "Does Salesforce support CSV exports?",
            self.docs,
        )

        self.assertEqual(result.status, "abstain")
        self.assertIsNone(result.passage)
        self.assertEqual(result.message, "I don't have an answer for that.")

    def test_single_generic_term_is_not_enough_evidence(self):
        result = answer_question("password", self.docs)

        self.assertEqual(result.status, "abstain")
        self.assertIsNotNone(result.message)

    def test_question_matching_multiple_settings_topics_requests_clarification(self):
        result = answer_question("What can I do in Settings?", self.docs)

        self.assertEqual(result.status, "ambiguous")
        self.assertGreaterEqual(len(result.alternatives), 2)

    def test_empty_question_abstains(self):
        self.assertEqual(answer_question("   ", self.docs).status, "abstain")

    def test_loader_reads_all_supported_files_recursively(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            (folder / "first.md").write_text(
                "# Password Help\nUse the reset link to change your password.",
                encoding="utf-8",
            )
            nested = folder / "nested"
            nested.mkdir()
            (nested / "D-22-billing.txt").write_text(
                "Monthly billing renews on your subscription date.",
                encoding="utf-8",
            )
            (nested / "api.json").write_text(
                json.dumps({
                    "docs": [{
                        "doc_id": "D-23",
                        "title": "API Authentication",
                        "content": "API requests require authentication keys.",
                    }],
                }),
                encoding="utf-8",
            )
            pdf_path = nested / "D-24-pdf-billing.pdf"
            write_text_pdf(pdf_path, "PDF billing renews monthly")

            word_path = nested / "D-25-word-export.docx"
            word = Document()
            word.core_properties.title = "Word Export Guide"
            word.add_paragraph("Word exports retain report filters.")
            table = word.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Maximum rows"
            table.cell(0, 1).text = "5000"
            word.save(word_path)

            docs = load_docs_from_folder(folder)

        self.assertEqual(len(docs), 5)
        self.assertEqual(
            {doc["title"] for doc in docs},
            {
                "Password Help",
                "D 22 Billing",
                "API Authentication",
                "D 24 Pdf Billing",
                "Word Export Guide",
            },
        )
        self.assertEqual(
            {doc["doc_id"] for doc in docs},
            {"first", "D-22", "D-23", "D-24", "D-25"},
        )
        result = answer_question("When does PDF billing renew monthly?", docs)
        self.assertEqual(result.status, "answer")
        self.assertEqual(result.doc_id, "D-24")
        result = answer_question("API requests require authentication", docs)
        self.assertEqual(result.status, "answer")
        self.assertEqual(result.doc_id, "D-23")
        result = answer_question("What do Word exports retain?", docs)
        self.assertEqual(result.status, "answer")
        self.assertEqual(result.doc_id, "D-25")
        self.assertIn("report filters", result.passage)
        self.assertIn("Maximum rows | 5000", next(
            doc["content"] for doc in docs if doc["doc_id"] == "D-25"
        ))


if __name__ == "__main__":
    unittest.main()
